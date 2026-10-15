"""PDF / DOCX / TXT extraction + pasted text normalization.

Security: extension allowlist + magic-byte sniff + size cap + control-char strip
+ length cap. Never returns raw stack traces — raises typed AppErrors instead.
"""
import io
import re
import unicodedata

from fastapi import UploadFile

from app.config import settings
from app.errors import DocumentError, EmptyDocument, FileTooLarge, UnsupportedFileType

_CONTROL_CHARS = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f\x7f]")
_WS = re.compile(r"[ \t]+")
_MULTI_NL = re.compile(r"\n{3,}")
# Renderers (PDF especially) emit invisible/formatting variants of real chars:
# soft hyphen must become "-" so IDs like API-240617 survive extraction.
_INVISIBLE = {"\u00ad": "-", "\u200b": "", "\u200c": "", "\u200d": "", "\ufeff": ""}


def _normalize(text: str) -> str:
    text = unicodedata.normalize("NFKC", text)
    for char, replacement in _INVISIBLE.items():
        text = text.replace(char, replacement)
    text = _CONTROL_CHARS.sub("", text)
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    text = "\n".join(_WS.sub(" ", line).rstrip() for line in text.split("\n"))
    text = _MULTI_NL.sub("\n\n", text)
    return text.strip()


def _cap(text: str) -> str:
    limit = settings.MAX_INPUT_CHARS
    if len(text) > limit:
        return text[:limit] + "\n\n[Document truncated]"
    return text


def _detect_source_type(filename: str | None) -> str:
    ext = (filename or "").rsplit(".", 1)[-1].lower() if filename and "." in filename else ""
    if ext not in settings.ALLOWED_EXTENSIONS:
        raise UnsupportedFileType(
            f"Unsupported file type. Allowed: {', '.join(sorted(settings.ALLOWED_EXTENSIONS))}."
        )
    return ext


def _read_capped(file: UploadFile) -> bytes:
    max_bytes = settings.max_upload_bytes
    chunks: list[bytes] = []
    total = 0
    while True:
        chunk = file.file.read(1024 * 1024)
        if not chunk:
            break
        total += len(chunk)
        if total > max_bytes:
            raise FileTooLarge(f"File exceeds the {settings.MAX_UPLOAD_MB} MB limit.")
        chunks.append(chunk)
    return b"".join(chunks)


def _check_magic(data: bytes, source_type: str) -> None:
    """Reject files whose content does not match their extension."""
    if source_type == "pdf" and not data.lstrip(b"\xef\xbb\xbf \t\n\r").startswith(b"%PDF"):
        raise UnsupportedFileType("File does not appear to be a valid PDF.")
    if source_type == "docx" and not data.startswith(b"PK"):
        raise UnsupportedFileType("File does not appear to be a valid DOCX archive.")


def _extract_pdf(data: bytes) -> str:
    try:
        import pymupdf

        with pymupdf.open(stream=data, filetype="pdf") as doc:
            parts = [page.get_text() for page in doc]
    except Exception as exc:
        raise DocumentError(f"Could not read PDF: {exc}") from exc
    return "\n".join(parts)


def _extract_docx(data: bytes) -> str:
    try:
        import docx  # python-docx

        document = docx.Document(io.BytesIO(data))
        parts = [p.text for p in document.paragraphs]
        for table in document.tables:
            for row in table.rows:
                parts.append(" | ".join(cell.text for cell in row.cells))
    except Exception as exc:
        raise DocumentError(f"Could not read DOCX: {exc}") from exc
    return "\n".join(parts)


def _decode_text(data: bytes) -> str:
    for encoding in ("utf-8-sig", "utf-8", "cp1252"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace")


def extract_upload(file: UploadFile) -> dict:
    source_type = _detect_source_type(file.filename)
    data = _read_capped(file)
    if not data:
        raise EmptyDocument("The uploaded file is empty.")
    _check_magic(data, source_type)

    if source_type == "pdf":
        raw = _extract_pdf(data)
    elif source_type == "docx":
        raw = _extract_docx(data)
    else:
        raw = _decode_text(data)

    text = _normalize(raw)
    if not text:
        raise EmptyDocument("No readable text could be extracted from this document.")
    return {
        "text": _cap(text),
        "source_type": source_type,
        "filename": file.filename,
    }


def normalize_pasted(text: str) -> str:
    """Public normalizer for already-incoming text (pasted/email/pipeline input)."""
    return _cap(_normalize(text))


def extract_text_content(text: str, source_hint: str | None = None) -> dict:
    """Normalize pasted text / email bodies (bypasses file handling)."""
    normalized = _normalize(_decode_text(text.encode("utf-8", errors="ignore")))
    if not normalized:
        raise EmptyDocument("No readable text was provided.")
    source_type = source_hint if source_hint in {"text", "email"} else _guess_source_type(normalized)
    return {
        "text": _cap(normalized),
        "source_type": source_type,
        "filename": None,
    }


def _guess_source_type(text: str) -> str:
    head = text[:400].lower()
    if head.startswith("subject:") or "\nsubject:" in head or head.startswith("from:"):
        return "email"
    return "text"
