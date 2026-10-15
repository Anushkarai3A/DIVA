"""Generate sample-deviation.pdf and sample-deviation.docx from sample-deviation.txt.

Run:  python sample-data/generate_samples.py   (backend venv: has PyMuPDF + python-docx)
"""
from pathlib import Path

HERE = Path(__file__).parent
SOURCE = HERE / "sample-deviation.txt"


def _unicode_font() -> dict:
    """Windows/Linux TTFs so en-dashes and °C survive extraction."""
    candidates = [
        Path("C:/Windows/Fonts/arial.ttf"),
        Path("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"),
        Path("/usr/share/fonts/truetype/liberation/LiberationSans-Regular.ttf"),
    ]
    for path in candidates:
        if path.exists():
            return {"fontname": "sans", "fontfile": str(path)}
    return {"fontname": "helv"}


def build_pdf(text: str, target: Path) -> None:
    import pymupdf

    font = _unicode_font()
    for fontsize in (11, 9, 7):  # retry smaller until the sample fits one page
        doc = pymupdf.open()
        page = doc.new_page(width=595, height=842)  # A4
        rect = pymupdf.Rect(56, 56, 539, 786)
        remaining = page.insert_textbox(
            rect, text, fontsize=fontsize, align=0, **font
        )
        if remaining >= 0:
            doc.save(target)
            doc.close()
            return
        doc.close()
    raise RuntimeError("sample text does not fit a single page even at 7pt")


def build_docx(text: str, target: Path) -> None:
    import docx

    document = docx.Document()
    for i, line in enumerate(text.split("\n")):
        if i == 0:
            document.add_heading(line, level=1)
        elif not line.strip():
            continue
        elif line.endswith(":") and len(line) < 60:
            document.add_heading(line, level=2)
        else:
            document.add_paragraph(line)
    document.save(target)


def main() -> None:
    text = SOURCE.read_text(encoding="utf-8")
    build_pdf(text, HERE / "sample-deviation.pdf")
    build_docx(text, HERE / "sample-deviation.docx")
    print(f"generated: {HERE / 'sample-deviation.pdf'}")
    print(f"generated: {HERE / 'sample-deviation.docx'}")


if __name__ == "__main__":
    main()
